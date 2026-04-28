import os
import glob
import datetime
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_report(results_dir):
    xml_files = glob.glob(os.path.join(results_dir, "junit_results.xml"))
    if not xml_files:
        return None

    wb = Workbook()
    
    # ---- Styles ----
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    light_blue_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)
    wrap_align = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "Execution Summary"
    
    ws_summary.merge_cells("A1:E1")
    ws_summary["A1"] = "Agrim Seller App - Test Execution Report"
    ws_summary["A1"].font = Font(bold=True, size=16, color="2F5496")
    ws_summary["A1"].alignment = Alignment(horizontal="center")

    # Parse results
    total = passed = failed = skipped = duration = 0
    for xml_file in xml_files:
        tree = ET.parse(xml_file)
        root = tree.getroot()
        for tc in root.iter("testcase"):
            total += 1
            duration += float(tc.get("time", 0))
            if tc.find("failure") is not None or tc.find("error") is not None:
                failed += 1
            elif tc.find("skipped") is not None:
                skipped += 1
            else:
                passed += 1

    stats = [
        ("Run Date", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total Tests", total),
        ("Passed", passed),
        ("Failed", failed),
        ("Skipped", skipped),
        ("Pass Rate", f"{(passed/total*100):.1f}%" if total > 0 else "0%"),
        ("Total Duration", f"{duration:.2f}s")
    ]
    
    for i, (label, val) in enumerate(stats, 3):
        ws_summary.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws_summary.cell(row=i, column=2, value=val)

    # Sheet 2: Details
    ws = wb.create_sheet("Test Case Details")
    headers = ["TC ID", "Module", "Test Name", "Status", "Duration (s)", "Error Details"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    row = 2
    for xml_file in xml_files:
        tree = ET.parse(xml_file)
        root = tree.getroot()
        for testcase in root.iter("testcase"):
            name = testcase.get("name", "Unknown")
            classname = testcase.get("classname", "Unknown")
            time = testcase.get("time", "0")
            
            failure = testcase.find("failure")
            error_el = testcase.find("error")
            
            if failure is not None: status = "FAIL"; msg = failure.text[:500] if failure.text else ""
            elif error_el is not None: status = "ERROR"; msg = error_el.text[:500] if error_el.text else ""
            else: status = "PASS"; msg = ""

            ws.cell(row=row, column=1, value=name.split("[")[0])
            ws.cell(row=row, column=2, value=classname.split(".")[-1])
            ws.cell(row=row, column=3, value=name)
            status_cell = ws.cell(row=row, column=4, value=status)
            status_cell.fill = pass_fill if status == "PASS" else fail_fill
            ws.cell(row=row, column=5, value=float(time))
            ws.cell(row=row, column=6, value=msg).alignment = wrap_align
            row += 1

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 80

    report_name = f"TestReport_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    report_path = os.path.join(results_dir, report_name)
    wb.save(report_path)
    return report_path
